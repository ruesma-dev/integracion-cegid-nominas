# application/use_cases/read_worker_convenio_from_templates_use_case.py
from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Optional

from openpyxl import load_workbook

from application.services.convenio_calendar import ConvenioCalendar
from config.config import Config

log = logging.getLogger(__name__)


@dataclass(frozen=True)
class WorkerConvenioMap:
    dni_to_ccc: Dict[str, str]
    dni_to_convenio: Dict[str, str]


class ReadWorkerConvenioFromTemplatesUseCase:
    def __init__(self, *, config: Config, convenio_calendar: ConvenioCalendar) -> None:
        self._config = config
        self._calendar = convenio_calendar

    def __call__(self) -> WorkerConvenioMap:
        tpl_cfg = self._config.join_config.templates or {}
        input_dir = Path(str(tpl_cfg.get("input_dir", "input")))
        pattern = str(tpl_cfg.get("pattern", "*.xlsx"))
        sheet_name = str(tpl_cfg.get("sheet", "") or "")

        templates = sorted(input_dir.glob(pattern))
        if not templates:
            raise FileNotFoundError(f"No se encontraron excels en {input_dir} con patrón '{pattern}'")

        dni_to_ccc: Dict[str, str] = {}
        dni_to_convenio: Dict[str, str] = {}

        for path in templates:
            wb = load_workbook(path, data_only=True)
            ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

            header_row = self._find_header_row(ws)
            if header_row is None:
                log.warning("No se encontró cabecera NIF en %s; omitido.", path.name)
                continue

            # columnas:
            # - B contiene el código compuesto CIF#CCC#CEGID#...
            # - la columna de NIF se detecta por la cabecera
            col_nif = self._find_col(ws, header_row, "NIF")
            if col_nif is None:
                log.warning("No se encontró columna NIF en %s; omitido.", path.name)
                continue

            for r in range(header_row + 1, ws.max_row + 1):
                nif_val = ws.cell(row=r, column=col_nif).value
                dni = self._norm(nif_val)
                if not dni:
                    continue

                code_b = ws.cell(row=r, column=2).value  # columna B
                ccc = self._extract_ccc(code_b)
                if not ccc:
                    continue

                dni_to_ccc[dni] = ccc
                dni_to_convenio[dni] = self._calendar.convenio_for_ccc(ccc)

        log.info("Mapeo DNI->CCC cargado: %s DNIs.", len(dni_to_ccc))
        return WorkerConvenioMap(dni_to_ccc=dni_to_ccc, dni_to_convenio=dni_to_convenio)

    @staticmethod
    def _find_header_row(ws) -> Optional[int]:
        max_scan = min(ws.max_row, 200)
        for r in range(1, max_scan + 1):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(row=r, column=c).value
                if v is None:
                    continue
                if str(v).strip() == "NIF" or str(v).strip() == "NIF":
                    return r
        return None

    @staticmethod
    def _find_col(ws, header_row: int, header_name: str) -> Optional[int]:
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=header_row, column=c).value
            if v is None:
                continue
            if str(v).strip() == header_name:
                return c
        return None

    @staticmethod
    def _extract_ccc(value: object) -> Optional[str]:
        if value is None:
            return None
        s = str(value).strip()
        if "#" not in s:
            return None
        parts = s.split("#")
        if len(parts) < 2:
            return None
        ccc = parts[1].strip()
        return ccc or None

    @staticmethod
    def _norm(value: object) -> Optional[str]:
        if value is None:
            return None
        s = str(value).strip().upper()
        return s if s else None
